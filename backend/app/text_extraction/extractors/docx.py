from typing import BinaryIO

from docx import Document as DocxDocument

from app.text_extraction.port import ExtractionResult


class DocxExtractor:
    def extract(self, stream: BinaryIO) -> ExtractionResult:
        doc = DocxDocument(stream)
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        # Tables carry real content in this document set (contracts, specs)
        # that paragraph text alone would miss.
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        text = "\n".join(paragraphs)
        return ExtractionResult(text=text, char_count=len(text))
