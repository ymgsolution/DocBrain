from typing import BinaryIO

from pptx import Presentation

from app.text_extraction.port import ExtractionResult


class PptxExtractor:
    def extract(self, stream: BinaryIO) -> ExtractionResult:
        presentation = Presentation(stream)
        parts: list[str] = []
        for index, slide in enumerate(presentation.slides, start=1):
            texts = [
                shape.text_frame.text
                for shape in slide.shapes
                if shape.has_text_frame and shape.text_frame.text.strip()
            ]
            if texts:
                parts.append(f"# Slide {index}\n" + "\n".join(texts))
        text = "\n\n".join(parts)
        return ExtractionResult(text=text, char_count=len(text))
