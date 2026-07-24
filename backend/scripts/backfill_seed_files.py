"""One-time backfill: write real, valid, minimal file content for every
document_version whose storage_path has no file on disk — scripts/seed.py
fabricates realistic documents/document_versions rows (filenames, checksums,
sizes) but never wrote actual bytes for the originally-seeded corpus (see
PROJECT_STATUS.md §10). Content is generated from each version's real
title/description/change-note, so it's at least genuinely related to the
metadata rather than lorem ipsum, and checksum_sha256/size_bytes are
recomputed to match what's actually written, so the DB and disk agree.

Safe to re-run: anything that already has a file on disk is left untouched.

Usage: uv run python -m scripts.backfill_seed_files
"""

import io
import textwrap

from docx import Document as DocxDocument
from openpyxl import Workbook
from pptx import Presentation
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from sqlalchemy import select

from app.db.models import Document, DocumentVersion
from app.db.session import SessionLocal
from app.storage.checksum import sha256_of_stream
from app.storage.local_adapter import LocalFileSystemStorage


def _generate_docx(title: str, body: str) -> bytes:
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    doc.add_paragraph(body)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _generate_xlsx(title: str, body: str) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Sheet1"
    sheet.append([title])
    sheet.append([body])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _generate_pptx(title: str, body: str) -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = title
    slide.placeholders[1].text = body
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _generate_pdf(title: str, body: str) -> bytes:
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    c.setFont("Helvetica-Bold", 16)
    c.drawString(72, 720, title[:90])
    c.setFont("Helvetica", 11)
    text = c.beginText(72, 690)
    for line in textwrap.wrap(body, 95) or [""]:
        text.textLine(line)
    c.drawText(text)
    c.save()
    return buffer.getvalue()


def _generate_plain(title: str, body: str) -> bytes:
    return f"{title}\n\n{body}\n".encode()


_GENERATORS = {
    "docx": _generate_docx,
    "xlsx": _generate_xlsx,
    "pptx": _generate_pptx,
    "pdf": _generate_pdf,
    "txt": _generate_plain,
    "md": _generate_plain,
    "csv": _generate_plain,
}


def main() -> None:
    db = SessionLocal()
    storage = LocalFileSystemStorage()
    try:
        versions = list(db.scalars(select(DocumentVersion)))
        written = 0
        already_present = 0
        unsupported = 0

        for version in versions:
            try:
                with storage.open_for_read(version.storage_path):
                    pass
                already_present += 1
                continue
            except FileNotFoundError:
                pass

            extension = version.original_filename.rsplit(".", 1)[-1].lower()
            generator = _GENERATORS.get(extension)
            if generator is None:
                unsupported += 1
                continue

            document = db.get(Document, version.document_id)
            title = document.title if document else version.original_filename
            body = (
                (document.description if document and document.description else None)
                or version.change_note
                or "Seed fixture content generated during backfill."
            )
            content = generator(title, body)

            temp_path = storage.save_temp(io.BytesIO(content))
            storage.commit(temp_path, version.storage_path)

            version.size_bytes = len(content)
            version.checksum_sha256 = sha256_of_stream(io.BytesIO(content))
            written += 1

        db.commit()
        print(
            f"Wrote {written} fixture file(s), {already_present} already existed, "
            f"{unsupported} unsupported extension(s) skipped."
        )
    finally:
        db.close()


if __name__ == "__main__":
    main()
