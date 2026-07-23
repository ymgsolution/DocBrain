"""Each extractor is a pure bytes-in/text-out function, so these build a real
fixture in-memory with the same library that would produce it in the wild
(python-docx, openpyxl, python-pptx) rather than committing binary fixture
files. pypdf can't easily *write* real text, so the PDF case only proves the
extractor handles a real, valid PDF without crashing — pypdf's own read path
is a well-tested third-party concern, not something to re-verify here.
"""

import io

import openpyxl
from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfWriter

from app.text_extraction.extractors.docx import DocxExtractor
from app.text_extraction.extractors.pdf import PdfExtractor
from app.text_extraction.extractors.plain import PlainTextExtractor
from app.text_extraction.extractors.pptx import PptxExtractor
from app.text_extraction.extractors.xlsx import XlsxExtractor


def test_plain_extractor_reads_utf8_text() -> None:
    stream = io.BytesIO("Quarterly report — revenue up 12%.".encode())
    result = PlainTextExtractor().extract(stream)
    assert result.text == "Quarterly report — revenue up 12%."
    assert result.char_count == len(result.text)


def test_plain_extractor_does_not_raise_on_bad_bytes() -> None:
    stream = io.BytesIO(b"valid text \xff\xfe more text")
    result = PlainTextExtractor().extract(stream)
    assert "valid text" in result.text
    assert "more text" in result.text


def test_docx_extractor_reads_paragraphs_and_tables() -> None:
    doc = DocxDocument()
    doc.add_paragraph("Onboarding Guide")
    doc.add_paragraph("Complete these steps before day one.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Step"
    table.rows[0].cells[1].text = "Owner"
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    result = DocxExtractor().extract(buffer)
    assert "Onboarding Guide" in result.text
    assert "Complete these steps before day one." in result.text
    assert "Step | Owner" in result.text
    assert result.char_count == len(result.text)


def test_xlsx_extractor_reads_cell_values_per_sheet() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    assert sheet is not None  # a freshly created Workbook always has one
    sheet.title = "Q1"
    sheet.append(["Category", "Amount"])
    sheet.append(["Travel", 4200])
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)

    result = XlsxExtractor().extract(buffer)
    assert "# Q1" in result.text
    assert "Category | Amount" in result.text
    assert "Travel | 4200" in result.text


def test_pptx_extractor_reads_slide_text() -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Roadmap Review"
    buffer = io.BytesIO()
    presentation.save(buffer)
    buffer.seek(0)

    result = PptxExtractor().extract(buffer)
    assert "Slide 1" in result.text
    assert "Roadmap Review" in result.text


def test_pdf_extractor_handles_a_real_pdf_without_raising() -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    buffer.seek(0)

    result = PdfExtractor().extract(buffer)
    # A genuinely blank page has no text — 0 is the correct outcome here,
    # not a failure.
    assert result.char_count == 0
    assert result.text == ""
